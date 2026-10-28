"""
Merges the User's Original 19-Slide Master Presentation with the 6 Professor-Feedback Slides.
Preserves 100% of user's original content, photos, layouts, and styles.
Adds the missing professor feedback items without any duplication:
1. The 6 Prioritized Core Modules (Slide 9)
2. Complete 17-Stage End-to-End Operational Workflow (Slide 16)
3. Inventory Accounting Ledger: Mathematical Balancing Equation & Immutable Audit Schema (Slide 17)
4. Replenishment Engineering: Demand, Safety Stock, ROP Formula & JIT Kanban (Slide 18)
5. Comprehensive End-to-End Process FMEA Matrix (Slide 21)
6. Evaluation Phasing & Review 2 Working Chain 10-Step Demonstration Script (Slide 24)
Total: 25 slides.
"""

import os
import shutil
import win32com.client
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

USER_PPTX_PATH = r"C:\Users\piyush\.gemini\antigravity\brain\d9f1016b-493a-4753-9c61-62a02f0730b4\.user_uploaded\media_1793163075130_2502535c.pptx"
OUTPUT_PPTX_PATH = r"C:\Users\piyush\Documents\Pharmacy_Management_System\TQM_Pharm_Presentation_Piyush_Bhatt.pptx"
OUTPUT_PDF_PATH = r"C:\Users\piyush\Documents\Pharmacy_Management_System\TQM_Pharm_Presentation_Piyush_Bhatt.pdf"
ARTIFACT_DIR = r"C:\Users\piyush\.gemini\antigravity\brain\d9f1016b-493a-4753-9c61-62a02f0730b4"

def build_merged_presentation():
    print(f"Loading user master presentation from:\n{USER_PPTX_PATH}")
    prs = Presentation(USER_PPTX_PATH)
    blank_layout = prs.slide_layouts[6]
    orig_count = len(prs.slides)
    print(f"Original user slide count: {orig_count}")

    # Visual Theme Palette
    NAVY_DARK = RGBColor(10, 37, 64)       # #0A2540
    NAVY_CARD = RGBColor(18, 52, 86)       # #123456
    TEAL_ACCENT = RGBColor(0, 163, 137)    # #00A389
    TEAL_LIGHT = RGBColor(220, 245, 240)   # #DCF5F0
    BG_LIGHT = RGBColor(248, 250, 252)     # #F8FAFC
    CARD_BG = RGBColor(255, 255, 255)      # #FFFFFF
    BORDER_COLOR = RGBColor(226, 232, 240) # #E2E8F0
    TEXT_MAIN = RGBColor(30, 41, 59)       # #1E293B
    TEXT_WHITE = RGBColor(255, 255, 255)   # #FFFFFF
    AMBER_WARN = RGBColor(217, 119, 6)     # #D97706
    BLUE_INFO = RGBColor(2, 132, 199)      # #0284C7
    PURPLE_ACCENT = RGBColor(124, 58, 237) # #7C3AED
    CORAL_DANGER = RGBColor(225, 29, 72)   # #E11D48

    def set_bg(slide, color=BG_LIGHT):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category_text):
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(0.35))
        tf_c = cat_box.text_frame
        tf_c.word_wrap = True
        p_c = tf_c.paragraphs[0]
        p_c.text = category_text.upper()
        p_c.font.size = Pt(10.5)
        p_c.font.bold = True
        p_c.font.color.rgb = TEAL_ACCENT
        p_c.font.name = "Segoe UI"

        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.733), Inches(0.8))
        tf_t = title_box.text_frame
        tf_t.word_wrap = True
        p_t = tf_t.paragraphs[0]
        p_t.text = title_text
        p_t.font.size = Pt(21)
        p_t.font.bold = True
        p_t.font.color.rgb = NAVY_DARK
        p_t.font.name = "Segoe UI"

    def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=BORDER_COLOR):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1.2)
        return shape

    def add_stacked_cards(slide, title, category, cards_data):
        set_bg(slide)
        add_header(slide, title, category)
        n = len(cards_data)
        total_h = 5.2
        gap = 0.14
        h = (total_h - gap * (n - 1)) / n
        y = 1.75
        for heading, body, *opt in cards_data:
            col = opt[0] if opt else TEAL_ACCENT
            c = add_card(slide, Inches(0.8), Inches(y), Inches(11.733), Inches(h))
            bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(y), Inches(0.12), Inches(h))
            bar.fill.solid()
            bar.fill.fore_color.rgb = col
            bar.line.fill.background()

            tb = slide.shapes.add_textbox(Inches(1.08), Inches(y + 0.05), Inches(11.25), Inches(h - 0.1))
            tf = tb.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = heading
            p.font.size = Pt(11.5 if n <= 4 else 10.5)
            p.font.bold = True
            p.font.color.rgb = NAVY_DARK
            p.space_after = Pt(2)

            for line in body.split('\n'):
                p2 = tf.add_paragraph()
                p2.text = line
                p2.font.size = Pt(10 if n <= 4 else 9)
                p2.font.color.rgb = TEXT_MAIN
                p2.space_after = Pt(1)
            y += h + gap

    def add_table_slide(slide, title, category, headers, rows_data, col_widths=None):
        set_bg(slide)
        add_header(slide, title, category)
        card_w = Inches(11.733)
        card_h = Inches(5.15)
        add_card(slide, Inches(0.8), Inches(1.75), card_w, card_h)

        table_shape = slide.shapes.add_table(len(rows_data) + 1, len(headers), Inches(0.95), Inches(1.9), Inches(11.433), Inches(4.8))
        tbl = table_shape.table

        if col_widths:
            for idx, w in enumerate(col_widths):
                tbl.columns[idx].width = Inches(w)

        for idx, h_text in enumerate(headers):
            cell = tbl.cell(0, idx)
            cell.text = h_text
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY_DARK
            p = cell.text_frame.paragraphs[0]
            p.font.bold = True
            p.font.size = Pt(10)
            p.font.color.rgb = TEXT_WHITE

        for r_idx, row in enumerate(rows_data, start=1):
            bg = BG_LIGHT if r_idx % 2 == 1 else CARD_BG
            for c_idx, val in enumerate(row):
                cell = tbl.cell(r_idx, c_idx)
                cell.text = str(val)
                cell.fill.solid()
                cell.fill.fore_color.rgb = bg
                p = cell.text_frame.paragraphs[0]
                p.font.size = Pt(9)
                p.font.color.rgb = TEXT_MAIN

    # =========================================================================
    # CREATE NEW SLIDE 1: 6 Prioritized Core Modules
    # =========================================================================
    new_s1 = prs.slides.add_slide(blank_layout)
    set_bg(new_s1)
    add_header(new_s1, "System Architecture: The 6 Prioritized Core Modules", "RUBRIC: MODULES & FEATURES PRIORITIZATION")

    mod_cards = [
        ("1. Product & Medicine Master", "• SKU, brand & generic salt master.\n• Therapeutic category & rack location.\n• Schedule H / X narcotic flag.\n• Cold-chain (2-8°C) boolean flag.\n• MRP, dealer cost, and GST %.\n• Drug interaction rule definitions.", TEAL_ACCENT),
        ("2. Supplier & Purchase/GRN", "• Approved vendor registry & terms.\n• Automated PO generation based on ROP.\n• Inward Goods Receipt Note (GRN).\n• Pay-on-Sale consignment inwarding.\n• Zero upfront dealer capital lock-in.\n• Discrepancy & damage logging.", BLUE_INFO),
        ("3. Inventory & FEFO Ledger", "• Multi-batch inventory tracking.\n• Automated FEFO picking sequence.\n• Visual countdown badges (<30d, <90d).\n• Perpetual inventory balance ledger.\n• Storage bin & rack allocation.\n• Cold-chain IoT fridge telemetry.", AMBER_WARN),
        ("4. Prescription & CDSS Shield", "• Digital eRx intake & patient linkage.\n• CDSS allergy cross-check engine.\n• Toxic drug-drug interaction shield.\n• Shingo Poka-Yoke checkout lockout.\n• Schedule X dual-PIN narcotic record.\n• Doctor MCI/DMC registration log.", CORAL_DANGER),
        ("5. POS & Multi-Tender Billing", "• High-speed barcode catalog search.\n• 1-Click approved eRx import to cart.\n• Split tender (Cash, UPI QR, Card).\n• Instant change return calculator.\n• Official 80mm GST tax invoice print.\n• Customer return & refund credit.", TEAL_ACCENT),
        ("6. Replenishment & Delivery", "• Reorder Point (ROP) dynamic formula.\n• Two-Bin Kanban buffer execution.\n• Pay-on-Sale supplier settlement.\n• Dispatch queue for delivery fleet.\n• Insulated cold-bag checklist.\n• 4-Digit SMS OTP proof of delivery.", PURPLE_ACCENT)
    ]
    col_w2 = Inches(3.64)
    for idx, (title, body, col) in enumerate(mod_cards):
        col_idx = idx % 3
        row_idx = idx // 3
        x = Inches(0.8) + Inches(col_idx * 4.04)
        y = Inches(1.8) + Inches(row_idx * 2.65)
        add_card(new_s1, x, y, col_w2, Inches(2.45))
        bar = new_s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, col_w2, Inches(0.4))
        bar.fill.solid()
        bar.fill.fore_color.rgb = col
        bar.line.fill.background()

        tb = new_s1.shapes.add_textbox(x, y + Inches(0.04), col_w2, Inches(0.35))
        tf = tb.text_frame
        p = tf.paragraphs[0]
        p.text = title
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(10.5)
        p.font.bold = True
        p.font.color.rgb = TEXT_WHITE

        tb2 = new_s1.shapes.add_textbox(x + Inches(0.12), y + Inches(0.45), col_w2 - Inches(0.24), Inches(1.95))
        tf2 = tb2.text_frame
        tf2.word_wrap = True
        for line in body.split('\n'):
            p = tf2.add_paragraph() if tf2.paragraphs[0].text else tf2.paragraphs[0]
            p.text = line
            p.font.size = Pt(9)
            p.font.color.rgb = TEXT_MAIN
            p.space_after = Pt(1.5)

    # =========================================================================
    # CREATE NEW SLIDE 2: 17-Step Complete Operational Workflow
    # =========================================================================
    new_s2 = prs.slides.add_slide(blank_layout)
    add_stacked_cards(new_s2, "End-to-End Pharmacy Workflow: Complete 17-Stage Operational Chain", "RUBRIC: PROCESS COMPLETENESS & MAIN SYSTEM WORKFLOW", [
        ("Phase I: Inbound Procurement & Quality Receiving (Steps 1 to 5)",
         "1. Wholesaler / Supplier Master -> 2. Purchase Planning (Demand Velocity & Lead Time) -> 3. Purchase Order (PO) Generation -> 4. Goods Receipt Note (GRN) Inwarding -> 5. Physical Inspection & Expiry Date Verification (<180d rejected). Stock enters Pay-on-Sale consignment ledger at ₹0 upfront capital.", TEAL_ACCENT),
        ("Phase II: Warehouse Storage & Lean Replenishment (Steps 6 to 8)",
         "6. Master Inventory Putaway (Rack allocation) -> 7. Replenishment Execution (Dynamic Reorder Point ROP & Two-Bin Kanban) -> 8. IoT Cold-Chain Monitoring (Continuous 2.0°C–8.0°C telemetry with acoustic excursion alarms).", BLUE_INFO),
        ("Phase III: Clinical Verification & POS Dispensing (Steps 9 to 13)",
         "9. Digital Prescription Intake (eRx) -> 10. Clinical Pharmacist Verification -> 11. CDSS Allergy & Drug Interaction Screening (Shingo Poka-Yoke hard-stop) -> 12. FEFO Algorithmic Dispensing -> 13. POS Checkout & Split Tender (Cash change calculator, Dynamic UPI QR, Card).", AMBER_WARN),
        ("Phase IV: Last-Mile Delivery, Returns & Governance (Steps 14 to 17)",
         "14. Home Delivery Fleet Dispatch & 4-Digit SMS OTP Handover -> 15. Customer Returns / Supplier Consignment Settlement -> 16. Real-Time Inventory Ledger Update -> 17. Daily Z-Report Cash Reconciliation & Immutable Audit Trail Logging.", PURPLE_ACCENT)
    ])

    # =========================================================================
    # CREATE NEW SLIDE 3: Inventory Accounting Ledger
    # =========================================================================
    new_s3 = prs.slides.add_slide(blank_layout)
    set_bg(new_s3)
    add_header(new_s3, "Inventory Ledger: Mathematical Balancing Equation & Immutable Audit", "RUBRIC: INVENTORY ACCOUNTING & TRACEABILITY ENGINE")

    add_card(new_s3, Inches(0.8), Inches(1.8), Inches(11.733), Inches(1.5), bg_color=NAVY_CARD, border_color=TEAL_ACCENT)
    tb_eq = new_s3.shapes.add_textbox(Inches(1.0), Inches(1.85), Inches(11.333), Inches(1.4))
    tf_eq = tb_eq.text_frame
    tf_eq.word_wrap = True
    pe0 = tf_eq.paragraphs[0]
    pe0.text = "MATHEMATICAL PERPETUAL INVENTORY BALANCING EQUATION"
    pe0.font.size = Pt(11)
    pe0.font.bold = True
    pe0.font.color.rgb = TEAL_ACCENT
    pe0.space_after = Pt(4)

    pe1 = tf_eq.add_paragraph()
    pe1.text = "Closing Stock = Opening + Receipts (GRN) + Customer Returns − Sales (POS) − Supplier Returns − Damage/Expired − Transfers Out + Transfers In ± Audit Adjustments"
    pe1.font.size = Pt(13)
    pe1.font.bold = True
    pe1.font.color.rgb = TEXT_WHITE
    pe1.space_after = Pt(4)

    pe2 = tf_eq.add_paragraph()
    pe2.text = "Every transaction creates a perpetual double-entry journal record. Naive scalar 'stock = stock - 1' is strictly prohibited."
    pe2.font.size = Pt(9.5)
    pe2.font.color.rgb = TEAL_LIGHT

    card_w_s = Inches(5.72)
    cards14_sub = [
        ("Immutable Ledger Data Schema", 
         "Every inventory movement writes an audit ledger entry containing:\n\n• Transaction ID: Unique UUID (e.g. TX-INV-2026-09812).\n• Timestamp: ISO 8601 millisecond-precision UTC date/time.\n• User ID & Role: Actor executing transaction (e.g. PBHATT, Cashier).\n• SKU & Batch Code: Physical batch identifier (e.g. DL-2026-B1).\n• Action Type: GRN_INWARD, POS_SALE, CUST_RETURN, WRITE_OFF.\n• Old Value -> Delta -> New Value: Exact numerical quantity change.\n• Reason & Reference Doc: Linked PO ID, Invoice ID, or Audit Note.", Inches(0.8), TEAL_ACCENT),
        ("Accounting vs Naive Retailing", 
         "Why TQM Requires a Mathematical Inventory Ledger:\n\n• Fraud & Shrinkage Interception: Detects unrecorded pilferage, unauthorized till adjustments, or ghost stock.\n• Lot-to-Lot Financial Reconciliation: Enables exact Cost of Goods Sold (COGS) calculation under FIFO/FEFO accounting.\n• Strict Regulatory Defense: Satisfies statutory drug inspector audits requiring proof of origin and chain of custody for every pill.\n• Clean Z-Report Alignment: Matches physical shelf count to system expected count during daily register closure.", Inches(6.813), BLUE_INFO)
    ]
    for title, body, left, col in cards14_sub:
        add_card(new_s3, left, Inches(3.45), card_w_s, Inches(3.45))
        bar = new_s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, Inches(3.45), card_w_s, Inches(0.4))
        bar.fill.solid()
        bar.fill.fore_color.rgb = col
        bar.line.fill.background()

        tb = new_s3.shapes.add_textbox(left, Inches(3.49), card_w_s, Inches(0.35))
        tf = tb.text_frame
        p = tf.paragraphs[0]
        p.text = title
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = TEXT_WHITE

        tb2 = new_s3.shapes.add_textbox(left + Inches(0.15), Inches(3.95), card_w_s - Inches(0.3), Inches(2.85))
        tf2 = tb2.text_frame
        tf2.word_wrap = True
        for line in body.split('\n'):
            p = tf2.add_paragraph() if tf2.paragraphs[0].text else tf2.paragraphs[0]
            p.text = line
            p.font.size = Pt(9.5)
            p.font.color.rgb = TEXT_MAIN
            p.space_after = Pt(2)

    # =========================================================================
    # CREATE NEW SLIDE 4: Replenishment Engineering & JIT Kanban
    # =========================================================================
    new_s4 = prs.slides.add_slide(blank_layout)
    set_bg(new_s4)
    add_header(new_s4, "Replenishment Engineering: Demand, Safety Stock, ROP & JIT Kanban", "RUBRIC: LEAN SUPPLY CHAIN & JIT REPLENISHMENT")

    add_card(new_s4, Inches(0.8), Inches(1.8), Inches(11.733), Inches(1.5), bg_color=NAVY_CARD, border_color=TEAL_ACCENT)
    tb_rop = new_s4.shapes.add_textbox(Inches(1.0), Inches(1.85), Inches(11.333), Inches(1.4))
    tf_rop = tb_rop.text_frame
    tf_rop.word_wrap = True
    pr0 = tf_rop.paragraphs[0]
    pr0.text = "MATHEMATICAL REPLENISHMENT MODEL (REORDER POINT & SAFETY STOCK)"
    pr0.font.size = Pt(11)
    pr0.font.bold = True
    pr0.font.color.rgb = TEAL_ACCENT
    pr0.space_after = Pt(4)

    pr1 = tf_rop.add_paragraph()
    pr1.text = "Reorder Point (ROP) = (Average Daily Demand × Lead Time) + Safety Stock\nSafety Stock (SS) = (Max Daily Demand × Max Lead Time) − (Avg Daily Demand × Avg Lead Time)"
    pr1.font.size = Pt(12.5)
    pr1.font.bold = True
    pr1.font.color.rgb = TEXT_WHITE
    pr1.space_after = Pt(4)

    pr2 = tf_rop.add_paragraph()
    pr2.text = "When Current Stock + On-Order <= ROP, the system triggers an automated purchase recommendation and generates a draft PO."
    pr2.font.size = Pt(9.5)
    pr2.font.color.rgb = TEAL_LIGHT

    cards15_sub = [
        ("JIT Pay-on-Sale Consignment Workflow", 
         "Lean Capital Management in Actual Pharmacy Practice:\n\n• Zero Upfront Working Capital: New stock enters inventory upon GRN receipt with ₹0 upfront payment to the pharmaceutical company.\n• Dynamic Sales Deductions: When a unit is sold at POS (e.g. Sold @ ₹100 MRP, Dealer Cost = ₹10), the system dynamically computes company liability (₹10) and records ₹90 gross pharmacy profit.\n• Dispute-Free Settlements: Automated supplier settlement vouchers generate payments strictly for verified sold quantities.\n• Unsold Stock Safety: Unsold batches remain supplier property or can be returned without financial loss.", Inches(0.8), TEAL_ACCENT),
        ("Two-Bin Kanban Physical Pull System", 
         "Visual Kanban Execution on the Dispensary Floor:\n\n• Bin 1 (Primary Active Dispensing): Kept at cashier POS counter for immediate patient fulfillment.\n• Bin 2 (Reserve Buffer Stock): Stored in backroom FEFO warehouse racks.\n• Kanban Trigger Mechanism: When Bin 1 is emptied, its physical Kanban barcode card is scanned. This initiates an immediate warehouse internal transfer to replenish Bin 1.\n• Upstream PO Automation: Scanning the Kanban card automatically checks current total stock against ROP to dispatch a digital purchase order to the wholesaler.", Inches(6.813), BLUE_INFO)
    ]
    for title, body, left, col in cards15_sub:
        add_card(new_s4, left, Inches(3.45), card_w_s, Inches(3.45))
        bar = new_s4.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, Inches(3.45), card_w_s, Inches(0.4))
        bar.fill.solid()
        bar.fill.fore_color.rgb = col
        bar.line.fill.background()

        tb = new_s4.shapes.add_textbox(left, Inches(3.49), card_w_s, Inches(0.35))
        tf = tb.text_frame
        p = tf.paragraphs[0]
        p.text = title
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = TEXT_WHITE

        tb2 = new_s4.shapes.add_textbox(left + Inches(0.15), Inches(3.95), card_w_s - Inches(0.3), Inches(2.85))
        tf2 = tb2.text_frame
        tf2.word_wrap = True
        for line in body.split('\n'):
            p = tf2.add_paragraph() if tf2.paragraphs[0].text else tf2.paragraphs[0]
            p.text = line
            p.font.size = Pt(9.5)
            p.font.color.rgb = TEXT_MAIN
            p.space_after = Pt(2)

    # =========================================================================
    # CREATE NEW SLIDE 5: Comprehensive End-to-End Process FMEA
    # =========================================================================
    new_s5 = prs.slides.add_slide(blank_layout)
    fmea_headers = ["Process Stage", "Potential Failure Mode", "Root Cause", "Sev (S)", "Occ (O)", "Det (D)", "RPN", "TQM Poka-Yoke Engineering Control"]
    fmea_rows = [
        ["1. Purchase Planning", "Stockout of critical cardiac drug", "Inaccurate demand forecast & lead time lag", "8", "5", "4", "160", "Automated ROP calculation trigger generating PO recommendations."],
        ["2. Receiving & GRN", "Inwarding near-expiry / damaged batch", "Superficial visual inspection during receipt", "9", "4", "4", "144", "Inward GRN system rejects any batch with < 180 days shelf life."],
        ["3. Inventory Storage", "Cold-chain insulin denaturation", "Chamber compressor failure or power cut", "9", "3", "3", "81", "24/7 IoT temperature telemetry with audible alarms and POS lockout."],
        ["4. Prescription Intake", "Doctor handwriting misinterpretation", "Ambiguous paper script abbreviations", "8", "6", "5", "240", "Digital eRx intake and mandatory EHR patient matching."],
        ["5. Clinical Check", "Penicillin allergy anaphylaxis", "Pharmacist misses allergy history", "10", "4", "5", "200", "CDSS Poka-Yoke hard-stop shield blocking contra-indicated drugs."],
        ["6. Dispensing & Pack", "Look-alike sound-alike (LASA) swap", "Dispensary clutter and hurried picking", "8", "4", "4", "128", "Mandatory 2D barcode scan verification matching NDC before packing."],
        ["7. Billing & POS", "Cash drawer deficit / till imbalance", "Manual mental math change error", "4", "6", "3", "72", "Instant change return calculator + Dynamic UPI QR webhook."],
        ["8. Delivery & Returns", "Delivery to wrong address / fake return", "Unverified customer handover", "7", "4", "4", "112", "Mandatory 4-digit SMS OTP electronic Proof-of-Delivery (e-POD)."]
    ]
    add_table_slide(new_s5, "End-to-End Process Failure Mode & Effects Analysis (FMEA)", "RUBRIC: TQM RISK GOVERNANCE ACROSS COMPLETE LIFECYCLE", fmea_headers, fmea_rows, [1.5, 1.8, 1.8, 0.6, 0.6, 0.6, 0.7, 3.833])

    # =========================================================================
    # CREATE NEW SLIDE 6: Evaluation Phasing & Review 2 Demonstration Script
    # =========================================================================
    new_s6 = prs.slides.add_slide(blank_layout)
    set_bg(new_s6)
    add_header(new_s6, "Evaluation Phasing: Review 1 Milestones vs Review 2 Demonstration", "RUBRIC: ACADEMIC EVALUATION PHASING & WORKING CHAIN")

    cards24_sub = [
        ("Review 1 Deliverables (Foundations & Architecture)", 
         "100% Completed & Verified in Current Review:\n\n• Complete 17-Step Pharmacy Workflow Mapping (Supplier to Audit).\n• Stakeholder Identification & Layman-to-SRS Mapping Table.\n• Prioritization of the 6 Core Operational Modules.\n• Role-Based Access Control & 5 Role Process Flows.\n• Software Requirements Specification (IEEE 830 Standard).\n• Relational 14-Table Database Design with ACID Transactions.\n• Comprehensive 8-Stage Process FMEA Risk Matrix.\n• Working Dual Desktop (Tkinter) & Web (FastAPI) GUI.\n• 25 Automated Test Suites Passed (100% Reliability).", Inches(0.8), TEAL_ACCENT),
        ("Review 2 Horizon (The Working Chain Demonstration)", 
         "10-Step Live Verification Script for Review 2:\n\n1. Inward GRN from Wholesaler into Pay-on-Sale Consignment.\n2. Real-Time Balance Verification in Perpetual Inventory Ledger.\n3. Automated ROP Calculation & Draft PO Generation.\n4. Ingest eRx & Trigger CDSS Penicillin Allergy Poka-Yoke Interception.\n5. Cashier 1-Click Rx Import & FEFO Auto-Batch Selection.\n6. POS Split Tender Checkout (Cash change + Dynamic UPI QR).\n7. Dispatch Order to Delivery Fleet with 4-Digit SMS OTP Handover.\n8. Pay-on-Sale Supplier Settlement Voucher & Zero-Variance Z-Report.\n9. Complete End-of-Day Shift Close & Audit Log Review.\n10. Point-in-Time Database Snapshot Archival & Backup.", Inches(6.813), BLUE_INFO)
    ]
    for title, body, left, col in cards24_sub:
        add_card(new_s6, left, Inches(1.8), card_w_s, Inches(5.1))
        bar = new_s6.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, Inches(1.8), card_w_s, Inches(0.5))
        bar.fill.solid()
        bar.fill.fore_color.rgb = col
        bar.line.fill.background()

        tb = new_s6.shapes.add_textbox(left, Inches(1.86), card_w_s, Inches(0.4))
        tf = tb.text_frame
        p = tf.paragraphs[0]
        p.text = title
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = TEXT_WHITE

        tb2 = new_s6.shapes.add_textbox(left + Inches(0.15), Inches(2.45), card_w_s - Inches(0.3), Inches(4.3))
        tf2 = tb2.text_frame
        tf2.word_wrap = True
        for line in body.split('\n'):
            p = tf2.add_paragraph() if tf2.paragraphs[0].text else tf2.paragraphs[0]
            p.text = line
            p.font.size = Pt(9.5)
            p.font.color.rgb = TEXT_MAIN
            p.space_after = Pt(2.5)

    # =========================================================================
    # REORDER SLIDES TO INTERLEAVE SEAMLESSLY WITHOUT DUPLICATION
    # =========================================================================
    # The original 19 slides are at indices 0..18.
    # The 6 new slides were added at the end: indices 19, 20, 21, 22, 23, 24.
    # We want final sequence:
    # 0..7: Orig 1..8 (Title, Problem, Objectives, Stakeholders, Roles, Arch, SRS, Modules Overview)
    # 8: NEW 1 (The 6 Prioritized Core Modules)
    # 9..14: Orig 9..14 (Core Features, DB Design, ACID, CRUD, Initial UI, Main Workflow)
    # 15: NEW 2 (17-Stage Complete Operational Workflow)
    # 16: NEW 3 (Inventory Ledger Balancing Equation)
    # 17: NEW 4 (Replenishment Engineering ROP & Kanban)
    # 18..19: Orig 15..16 (Pareto & Fishbone, FMEA & PDCA)
    # 20: NEW 5 (Comprehensive Process FMEA Matrix)
    # 21..22: Orig 17..18 (Working App Setup, 5-Phase Demo)
    # 23: NEW 6 (Evaluation Phasing & Review 2 Demonstration Script)
    # 24: Orig 19 (Thank you slide)

    sldIdLst = prs.slides._sldIdLst
    # Get references to the 6 newly added sldId elements
    id_new6 = sldIdLst[24]
    id_new5 = sldIdLst[23]
    id_new4 = sldIdLst[22]
    id_new3 = sldIdLst[21]
    id_new2 = sldIdLst[20]
    id_new1 = sldIdLst[19]

    # Pop them from the end (in reverse order)
    sldIdLst.remove(id_new6)
    sldIdLst.remove(id_new5)
    sldIdLst.remove(id_new4)
    sldIdLst.remove(id_new3)
    sldIdLst.remove(id_new2)
    sldIdLst.remove(id_new1)

    # Now insert at exact positions:
    # 1. Insert new1 at index 8 (after orig slide 8)
    sldIdLst.insert(8, id_new1)

    # 2. Insert new2, new3, new4 after orig slide 14 (which is now at index 14)
    # So index 15, 16, 17
    sldIdLst.insert(15, id_new2)
    sldIdLst.insert(16, id_new3)
    sldIdLst.insert(17, id_new4)

    # 3. Insert new5 after orig slide 16 (which is now at index 19)
    # So index 20
    sldIdLst.insert(20, id_new5)

    # 4. Insert new6 after orig slide 18 (which is now at index 22)
    # So index 23
    sldIdLst.insert(23, id_new6)

    print(f"Final presentation slide count: {len(prs.slides)}")

    # Save merged presentation
    prs.save(OUTPUT_PPTX_PATH)
    print(f"[SUCCESS] Saved merged presentation to: {OUTPUT_PPTX_PATH}")

    # Copy to artifact directory
    artifact_pptx = os.path.join(ARTIFACT_DIR, "TQM_Pharm_Presentation_Piyush_Bhatt.pptx")
    shutil.copy2(OUTPUT_PPTX_PATH, artifact_pptx)
    print(f"[SUCCESS] Copied presentation to artifacts: {artifact_pptx}")

    # Export to PDF via PowerPoint COM
    try:
        ppt_app = win32com.client.Dispatch("PowerPoint.Application")
        deck = ppt_app.Presentations.Open(os.path.abspath(OUTPUT_PPTX_PATH), WithWindow=False)
        deck.SaveAs(os.path.abspath(OUTPUT_PDF_PATH), 32)
        deck.Close()
        ppt_app.Quit()
        print(f"[SUCCESS] Generated PDF successfully at: {OUTPUT_PDF_PATH}")
        artifact_pdf = os.path.join(ARTIFACT_DIR, "TQM_Pharm_Presentation_Piyush_Bhatt.pdf")
        shutil.copy2(OUTPUT_PDF_PATH, artifact_pdf)
        print(f"[SUCCESS] Copied PDF to artifacts: {artifact_pdf}")
    except Exception as e:
        print(f"[WARNING] PowerPoint COM export failed: {e}")

if __name__ == "__main__":
    build_merged_presentation()
